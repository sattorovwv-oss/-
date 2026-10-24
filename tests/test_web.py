from io import BytesIO
import uuid

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from tuman.config import Settings
from tuman.web import create_app
from .conftest import TEST_PASSWORD, csrf, login


def test_public_pages_and_security_headers(client):
    for path in ("/", "/login", "/register", "/auth/telegram", "/health"):
        response=client.get(path)
        assert response.status_code==200
        assert response.headers["x-frame-options"]=="DENY"
        assert "no-store" in response.headers["cache-control"]
    assert "httponly" in client.get("/login").headers.get("set-cookie", "").lower() or client.cookies.get("tuman_session")


def test_registration_never_creates_admin(client,store):
    token=csrf(client,"/register")
    response=client.post("/register",data={"csrf":token,"username":"fresh","name":"Client","password":TEST_PASSWORD,"confirm":TEST_PASSWORD,"role":"owner"},follow_redirects=False)
    assert response.status_code==303
    assert store.one("SELECT role FROM accounts WHERE username='fresh'")["role"]=="customer"
    assert client.get("/admin").status_code==403
    assert client.get("/admin/settings").status_code==403


@pytest.mark.parametrize("path",["/admin","/admin/scan","/admin/parcels","/admin/claims","/admin/clients","/admin/support","/admin/broadcast","/admin/settings","/profile"])
def test_admin_pages_render(client,path):
    login(client)
    assert client.get(path).status_code==200


def test_login_rotates_cookie_and_csrf(client):
    old=csrf(client,"/login")
    old_cookie=client.cookies.get("tuman_session")
    new=login(client)
    assert old!=new
    assert old_cookie!=client.cookies.get("tuman_session")


def test_csrf_and_role_protect_scanning(client,store):
    token=login(client,"china")
    data={"track":"TRACK123","warehouse":"china","request_id":str(uuid.uuid4())}
    assert client.post("/api/scans",json=data).status_code==403
    assert client.post("/api/scans",json={**data,"warehouse":"tajik"},headers={"X-CSRF-Token":token}).status_code==403
    assert store.parcel("TRACK123") is None
    assert client.post("/api/scans",json=data,headers={"X-CSRF-Token":token}).status_code==200


def test_customer_cannot_submit_scan_or_read_other_private_data(client,store,customer):
    store.scan(store.one("SELECT id FROM accounts WHERE role='china'")["id"],"china","TRACK123",str(uuid.uuid4()))
    token=login(client,"customer")
    assert client.post("/api/scans",json={"track":"TRACK123","warehouse":"china","request_id":str(uuid.uuid4())},headers={"X-CSRF-Token":token}).status_code==403
    response=client.get("/account/search?track=TRACK123")
    assert response.status_code==200
    assert "TRACK123" in response.text
    assert "Личные сведения" in response.text
    assert client.get("/admin/export").status_code==403


def test_admin_cannot_edit_other_warehouse(client,store):
    p=store.scan(store.one("SELECT id FROM accounts WHERE role='tajik'")["id"],"tajik","TRACK123",str(uuid.uuid4()))
    token=login(client,"china")
    response=client.post(f"/admin/parcels/{p['parcel_id']}",data={"csrf":token,"weight":"1","price":"20","stage":"30"})
    assert response.status_code==403


def test_telegram_exchange_enforces_origin_and_single_use(client,store):
    account=store.one("SELECT id FROM accounts WHERE role='china'")
    token=store.issue_link(account_id=account["id"])
    assert client.post("/api/auth/telegram",json={"token":token},headers={"Origin":"https://evil.example"}).status_code==403
    response=client.post("/api/auth/telegram",json={"token":token},headers={"Origin":"http://testserver"})
    assert response.status_code==200
    assert response.json()["redirect"]=="/admin"
    assert client.get("/admin").status_code==200
    assert client.post("/api/auth/telegram",json={"token":token},headers={"Origin":"http://testserver"}).status_code==400


def test_logout_revokes_session(client):
    token=login(client)
    cookie=client.cookies.get("tuman_session")
    assert client.post("/logout",data={"csrf":token},follow_redirects=False).status_code==303
    client.cookies.set("tuman_session",cookie)
    assert client.get("/admin",follow_redirects=False).status_code==303


def test_settings_are_owner_only(client):
    token=login(client,"china")
    assert client.get("/admin/settings").status_code==403
    assert client.post("/admin/settings",data={"csrf":token,"rate":"20","currency":"TJS"}).status_code==403


def test_xss_is_escaped_in_support(client,store,customer):
    store.create_ticket(customer["person_id"],'<script>alert("x")</script>')
    login(client)
    response=client.get("/admin/support")
    assert '&lt;script&gt;' in response.text
    assert '<script>alert("x")</script>' not in response.text


def test_excel_is_real_workbook_and_formula_safe(client,store):
    store.scan(store.one("SELECT id FROM accounts WHERE role='china'")["id"],"china","-ABC123",str(uuid.uuid4()))
    login(client,"china")
    response=client.get("/admin/export")
    assert response.status_code==200
    workbook=load_workbook(BytesIO(response.content))
    assert workbook.active.cell(2,1).value=="'-ABC123"
    assert workbook.active.cell(2,1).data_type=="s"


def test_production_cookie_is_secure(store):
    app=create_app(Settings(db_path=store.path,production=True,public_url="https://cargo.example"),store)
    with TestClient(app,base_url="https://cargo.example") as c:
        response=c.get("/login")
        assert "secure" in response.headers["set-cookie"].lower()
        assert "max-age=31536000" in response.headers["strict-transport-security"]


def test_too_large_request_is_rejected(client):
    assert client.post("/login",content="x"*65000).status_code==413
