"""Regressions reproduced from the first Windows run, not live Telegram calls."""
import asyncio
from datetime import datetime, timezone
from io import BytesIO
import sys
from unittest.mock import AsyncMock
import uuid
import zoneinfo

from aiogram import Bot
from aiogram.exceptions import TelegramBadRequest, TelegramNetworkError
from aiogram.methods import GetUpdates, SendMessage
from aiogram.types import CallbackQuery, Chat, Message, Update, User
from openpyxl import load_workbook
import pytest

from tuman.bot import build_dispatcher
from tuman.config import Settings
from tuman.network import RetryTelegramNetwork
from tuman.workers import deliver_once
from .conftest import login
from .test_bot import RecordingSession, feed


@pytest.mark.parametrize("url", [
    "http://localhost:8000", "https://localhost:8000", "https://test.localhost",
    "https://127.0.0.1:8000", "https://[::1]:8000", "https://192.168.1.1",
    "https://10.0.0.1", "https://warehouse.local", "https://warehouse",
    "https://cargo.example", "https://name:password@example.com", "http://example.com",
    "https://bad host.example.com", "https://cargo.example.com:99999",
    "https://cargo.example.com?key=value", "https://cargo.example.com/subpath",
])
def test_local_or_unsuitable_addresses_do_not_create_telegram_buttons(url):
    assert not Settings(public_url=url).has_public_web_url()


def test_public_https_address_can_create_telegram_buttons():
    assert Settings(public_url="https://cargo.example.com").has_public_web_url()


def test_local_mode_admin_commands_and_customer_website_do_not_emit_bad_urls(store):
    async def run():
        session = RecordingSession()
        bot = Bot("123456:TEST_ONLY_NOT_A_REAL_BOT_TOKEN", session=session)
        account = store.one("SELECT id FROM accounts WHERE role='owner'")
        with store.connection(True) as db:
            db.execute("INSERT INTO telegram_admins VALUES (101,?)", (account["id"],))
        store.telegram_person(101, "Client", language="ru")
        dp = build_dispatcher(store, Settings())
        await feed(dp, bot, "/panel")
        panel = session.sent[-1]
        assert "PUBLIC_URL" in panel.text
        buttons = [button for row in panel.reply_markup.inline_keyboard for button in row]
        assert not any(button.url for button in buttons)
        assert any(button.callback_data == "admin:stats" for button in buttons)
        assert any(button.callback_data == "admin:rate" for button in buttons)
        await feed(dp, bot, "/website")
        assert "Сайт пока не опубликован" in session.sent[-1].text
        assert session.sent[-1].reply_markup is None
        assert store.one("SELECT COUNT(*) n FROM login_tokens")["n"] == 0
        await dp.storage.close()
    asyncio.run(run())


def test_windows_without_timezone_database_can_render_customer_ticket_and_excel(client, store, customer, monkeypatch):
    original_path = zoneinfo.TZPATH
    zoneinfo.ZoneInfo.clear_cache()
    zoneinfo.reset_tzpath(())
    monkeypatch.setitem(sys.modules, "tzdata", None)
    try:
        with pytest.raises(zoneinfo.ZoneInfoNotFoundError):
            zoneinfo.ZoneInfo("Asia/Dushanbe")
        store.add_track(customer["person_id"], "WINDOWS123")
        admin_id = store.one("SELECT id FROM accounts WHERE role='owner'")["id"]
        store.decide_claim(store.one("SELECT id FROM claims")["id"], admin_id)
        store.create_ticket(customer["person_id"], "Проверка времени на Windows")
        login(client, "customer")
        response = client.get("/account")
        assert response.status_code == 200
        created = store.parcel("WINDOWS123")["created"]
        from tuman.web import DUSHANBE_TIMEZONE
        assert datetime.fromtimestamp(created, DUSHANBE_TIMEZONE).strftime("%d.%m.%Y %H:%M") in response.text
        login(client, "owner")
        response = client.get("/admin/export")
        assert response.status_code == 200
        sheet = load_workbook(BytesIO(response.content)).active
        assert sheet.cell(2, 7).value.endswith("+05:00")
        assert client.get("/admin/support").status_code == 200
    finally:
        zoneinfo.reset_tzpath(original_path)
        zoneinfo.ZoneInfo.clear_cache()


def test_transport_retries_transient_failure_but_not_telegram_api_rejections():
    async def run():
        request = SendMessage(chat_id=101, text="Test")
        retry = RetryTelegramNetwork(delay=0)
        transport = AsyncMock(side_effect=[TelegramNetworkError(method=request, message="[WinError 64]"), "accepted"])
        assert await retry(transport, None, request) == "accepted"
        assert transport.await_count == 2
        rejected = AsyncMock(side_effect=TelegramBadRequest(method=request, message="invalid request"))
        with pytest.raises(TelegramBadRequest):
            await retry(rejected, None, request)
        assert rejected.await_count == 1
        polling = GetUpdates()
        broken_polling = AsyncMock(side_effect=TelegramNetworkError(method=polling, message="network"))
        with pytest.raises(TelegramNetworkError):
            await retry(broken_polling, None, polling)
        assert broken_polling.await_count == 1  # Polling already owns its own backoff.
    asyncio.run(run())


class FlakyTelegramSession(RecordingSession):
    def __init__(self, failures):
        super().__init__()
        self.failures = failures
        self.send_attempts = 0
        self.middleware.register(RetryTelegramNetwork(delay=0))

    async def make_request(self, bot, method, timeout=None):
        if isinstance(method, SendMessage):
            self.send_attempts += 1
            if self.failures:
                self.failures -= 1
                raise TelegramNetworkError(method=method, message="[WinError 64] test failure; secret=DO_NOT_LOG_TRANSPORT_TEXT")
        return await super().make_request(bot, method, timeout)


def test_language_callback_recovers_after_windows_connection_reset(store):
    async def run():
        session = FlakyTelegramSession(failures=1)
        bot = Bot("123456:TEST_ONLY_NOT_A_REAL_BOT_TOKEN", session=session)
        dp = build_dispatcher(store, Settings())
        user = User(id=101, first_name="Client", is_bot=False)
        update = Update(update_id=201, callback_query=CallbackQuery(
            id="language-choice", from_user=user, chat_instance="test", data="lang:ru",
            message=Message(message_id=1, date=datetime.now(timezone.utc), chat=Chat(id=101, type="private")),
        ))
        await dp.feed_update(bot, update)
        assert "Добро пожаловать" in session.sent[-1].text
        assert session.send_attempts == 2
        assert store.one("SELECT lang FROM people WHERE telegram_id=101")["lang"] == "ru"
        assert store.one("SELECT COUNT(*) n FROM outbox")["n"] == 0
        await dp.storage.close()
    asyncio.run(run())


def test_persistent_connection_failure_keeps_business_data_and_queues_recovery(store, caplog):
    async def run():
        store.telegram_person(101, "Client", language="ru")
        session = FlakyTelegramSession(failures=3)
        bot = Bot("123456:TEST_ONLY_NOT_A_REAL_BOT_TOKEN", session=session)
        dp = build_dispatcher(store, Settings())
        await feed(dp, bot, "WINDOWS123")
        assert session.send_attempts == 3
        assert store.parcel("WINDOWS123") is not None
        assert store.one("SELECT COUNT(*) n FROM claims")["n"] == 1
        notice = store.one("SELECT * FROM outbox")
        assert notice["state"] == "pending" and notice["event_key"].startswith("connection:")
        # Duplicate error handling cannot enqueue the same recovery notice twice.
        update_id = notice["event_key"].rsplit(":", 1)[1]
        store.queue_connection_notice(101, update_id, "duplicate")
        assert store.one("SELECT COUNT(*) n FROM outbox")["n"] == 1
        await deliver_once(store, bot)
        assert store.one("SELECT state FROM outbox")["state"] == "sent"
        assert "/start" in session.sent[-1].text
        assert "DO_NOT_LOG_TRANSPORT_TEXT" not in caplog.text
        await dp.storage.close()
    asyncio.run(run())
